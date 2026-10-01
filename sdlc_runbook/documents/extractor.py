"""Text extraction from an uploaded requirements document (.txt, .md, .docx, .pdf).

The file type is checked from its content, not only its extension. Nothing is written
to disk; only the extracted text and a SHA-256 fingerprint leave this module. Documents
that are too large, encrypted, or have no extractable text (e.g. scanned PDFs) are
rejected with a clear message rather than silently truncated.
"""
import hashlib
import io
import zipfile
from dataclasses import dataclass
from pathlib import PurePath

from ..exceptions import DocumentError

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_CHARS = 200_000
MAX_PDF_PAGES = 300

_KINDS = {".txt": "text", ".md": "markdown", ".docx": "docx", ".pdf": "pdf"}


@dataclass(frozen=True)
class ExtractedDocument:
    filename: str
    kind: str
    sha256: str
    size_bytes: int
    pages: int | None
    text: str

    @property
    def chars(self) -> int:
        return len(self.text)


def _text_file(data: bytes) -> tuple[str, None]:
    if b"\x00" in data[:4096]:
        raise DocumentError("This file looks binary, not plain text.")
    try:
        return data.decode("utf-8-sig"), None
    except UnicodeDecodeError:
        try:
            return data.decode("cp1252"), None
        except UnicodeDecodeError as e:
            raise DocumentError("The text file is not UTF-8 encoded.") from e


def _docx(data: bytes) -> tuple[str, None]:
    if not data.startswith(b"PK\x03\x04"):
        raise DocumentError("This is not a valid .docx file.")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if "word/document.xml" not in z.namelist():
                raise DocumentError("This .docx has no document body.")
    except zipfile.BadZipFile as e:
        raise DocumentError("This .docx file is corrupted.") from e

    import docx  # python-docx

    doc = docx.Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts), None


def _pdf(data: bytes) -> tuple[str, int]:
    if not data.startswith(b"%PDF"):
        raise DocumentError("This is not a valid PDF file.")
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise DocumentError("This PDF is password-protected. Remove the protection and upload it again.")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise DocumentError(f"This PDF has {len(reader.pages)} pages; the limit is {MAX_PDF_PAGES}.")
        text = "\n\n".join((page.extract_text() or "") for page in reader.pages)
    except PdfReadError as e:
        raise DocumentError("This PDF could not be read.") from e
    return text, len(reader.pages)


def _normalise(text: str) -> str:
    lines = [line.rstrip() for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    out, blank = [], 0
    for line in lines:
        blank = blank + 1 if not line.strip() else 0
        if blank <= 1:
            out.append(line)
    return "\n".join(out).strip()


def extract_text(filename: str, data: bytes) -> ExtractedDocument:
    name = PurePath(filename or "document").name[:200]
    suffix = PurePath(name).suffix.lower()
    kind = _KINDS.get(suffix)
    if kind is None:
        raise DocumentError(f"Unsupported file type '{suffix or 'none'}'. Upload a .txt, .md, .docx or .pdf file.")
    if not data:
        raise DocumentError("The file is empty.")
    if len(data) > MAX_UPLOAD_BYTES:
        raise DocumentError(f"The file is {len(data) / 1_048_576:.1f} MB; the limit is {MAX_UPLOAD_BYTES // 1_048_576} MB.")

    reader = {"text": _text_file, "markdown": _text_file, "docx": _docx, "pdf": _pdf}[kind]
    raw, pages = reader(data)
    text = _normalise(raw)
    if not text:
        hint = " It may be a scanned document; upload a version with selectable text." if kind == "pdf" else ""
        raise DocumentError("No readable text was found in this document." + hint)
    if len(text) > MAX_CHARS:
        raise DocumentError(
            f"The document has {len(text):,} characters of text; the limit is {MAX_CHARS:,}. "
            "Upload the relevant section only."
        )
    return ExtractedDocument(name, kind, hashlib.sha256(data).hexdigest(), len(data), pages, text)
